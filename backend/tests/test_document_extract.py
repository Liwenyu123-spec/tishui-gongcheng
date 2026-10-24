from __future__ import annotations

from io import BytesIO


def test_docx_table_text_is_preserved():
    from docx import Document

    document = Document()
    document.add_paragraph("泵组阈值表")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "设备"
    table.cell(0, 1).text = "阈值"
    table.cell(0, 2).text = "单位"
    table.cell(1, 0).text = "二泵站"
    table.cell(1, 1).text = "105"
    table.cell(1, 2).text = "百分比"
    payload = BytesIO()
    document.save(payload)

    from app.services.document_extract import extract_text

    text, parser = extract_text("thresholds.docx", payload.getvalue())
    assert parser == "python-docx"
    assert "[表格 1]" in text
    assert "二泵站\t105\t百分比" in text
